"""Send the debts exported from Phenix to Alpha Group.

Reads only: the Excel file the accountant exports (customer, invoice number, date, amount,
remaining) and sends it to the app, which makes its debts agree with it. Needs Python 3 and
`pip install openpyxl`; nothing else.

    python phenix_sync.py check FILE.xlsx        read + validate the file, send nothing
    python phenix_sync.py dry   FILE.xlsx        ask the app what WOULD change (writes nothing)
    python phenix_sync.py apply FILE.xlsx        apply it (add --force for a held-back run)
    python phenix_sync.py undo  RUN_ID           put the latest applied run back as it was
    python phenix_sync.py runs                   the latest runs
    python phenix_sync.py watch                  watch the inbox folder and apply each new export

Settings: config.ini next to this file (see config.example.ini). The key can also come from
the ALPHA_SYNC_KEY environment variable.
"""
import configparser
import datetime as dt
import hashlib
import json
import os
import re
import shutil
import sys
import time
import urllib.error
import urllib.request

HERE = os.path.dirname(os.path.abspath(__file__))

# The export's columns by their header text; the data may sit one column to the left of its
# header (the program writes a blank "#" header), so the offset is detected, not assumed.
NEEDED = {
    "invoice": "رقم السند",
    "kind": "النوع",
    "customer": "العميل",
    "date": "التاريخ",
    "amount": "القيمة",
    "paid": "المسدد",
    "remaining": "الباقي",
}


class SyncError(Exception):
    pass


class FileBusy(SyncError):
    """The file is open in another program (Excel, or still being written): try again later."""


def _num(value):
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace(",", "").strip()
    try:
        return float(text)
    except ValueError:
        return None


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _date(value):
    if isinstance(value, (dt.datetime, dt.date)):
        return value.strftime("%Y-%m-%d")
    text = _text(value)
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", text)
    return f"{m.group(1)}-{m.group(2)}-{m.group(3)}" if m else text


def read_export(path):
    """-> (rows, ignored_returns, info). Raises SyncError when the file is not the debts export."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise SyncError("ثبّت المكتبة أولاً:  pip install openpyxl")
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except PermissionError:
        raise FileBusy("الملف مفتوح ببرنامج آخر (Excel أو قيد الكتابة)")
    except Exception as e:
        raise SyncError(f"تعذّر فتح الملف: {e}")
    ws = wb.worksheets[0]
    all_rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    all_rows = [r for r in all_rows if any(c not in (None, "") for c in r)]
    if len(all_rows) < 2:
        raise SyncError("الملف فارغ")

    header = [_text(c) for c in all_rows[0]]
    missing = [h for h in NEEDED.values() if h not in header]
    if missing:
        raise SyncError("الملف ليس تقرير الديون المتوقع، أعمدة مفقودة: " + "، ".join(missing))
    pos = {k: header.index(h) for k, h in NEEDED.items()}
    data = all_rows[1:]

    offset = None
    for off in (0, 1):  # data column = header column - off
        ok = 0
        sample = data[:30]
        for r in sample:
            try:
                cust = r[pos["customer"] - off]
                rem = r[pos["remaining"] - off]
                inv = r[pos["invoice"] - off]
                amt = r[pos["amount"] - off]
            except IndexError:
                break
            if _num(cust) is None and _text(cust) and _num(rem) is not None and _num(amt) is not None and re.fullmatch(r"\d+", _text(inv) or "x"):
                ok += 1
        if sample and ok == len(sample):
            offset = off
            break
    if offset is None:
        raise SyncError("ترتيب الأعمدة بالملف غير مطابق للمتوقع (اسم الزبون/الباقي/رقم السند). لم يُرسل شيء.")

    rows, returns, mismatches = [], 0, 0
    for r in data:
        def get(k):
            i = pos[k] - offset
            return r[i] if 0 <= i < len(r) else None

        invoice, customer = _text(get("invoice")), _text(get("customer"))
        amount, remaining, paid = _num(get("amount")), _num(get("remaining")), _num(get("paid"))
        if not invoice and not customer:
            continue
        if amount is None or remaining is None:
            raise SyncError(f"سطر بمبلغ غير مفهوم (رقم السند {invoice or '?'}). لم يُرسل شيء.")
        if amount <= 0 or "مرتجع" in _text(get("kind")):
            returns += 1
            continue
        if paid is not None and abs(amount - paid - remaining) > 1:
            mismatches += 1
        rows.append(
            {
                "invoice_number": invoice,
                "customer_name": customer,
                "amount": amount,
                "remaining": remaining,
                "invoice_date": _date(get("date")),
            }
        )
    if not rows:
        raise SyncError("لا توجد قوائم دين بالملف")
    if mismatches > max(3, len(rows) * 0.02):
        raise SyncError(f"{mismatches} سطر لا يطابق (القيمة − المسدد = الباقي)؛ الأعمدة غالباً غير صحيحة. لم يُرسل شيء.")
    info = {
        "rows": len(rows),
        "returns_ignored": returns,
        "customers": len({r["customer_name"] for r in rows}),
        "remaining_total": round(sum(r["remaining"] for r in rows), 2),
        "column_offset": offset,
    }
    return rows, returns, info


# ---------------------------------------------------------------- talking to the app
def load_config():
    cp = configparser.ConfigParser()
    cp.read(os.path.join(HERE, "config.ini"), encoding="utf-8-sig")  # tolerate the BOM Notepad adds
    cfg = {
        "url": cp.get("server", "url", fallback="").rstrip("/"),
        "key": os.environ.get("ALPHA_SYNC_KEY") or cp.get("server", "key", fallback=""),
        "inbox": cp.get("folders", "inbox", fallback=os.path.join(HERE, "inbox")),
        "interval": cp.getint("watch", "interval_seconds", fallback=30),
    }
    if not cfg["url"] or not cfg["key"]:
        raise SyncError("أكمل config.ini: عنوان الخادم (url) ومفتاح المزامنة (key)")
    return cfg


def call(cfg, method, path, body=None):
    req = urllib.request.Request(
        cfg["url"] + path,
        data=None if body is None else json.dumps(body).encode("utf-8"),
        method=method,
        headers={"Content-Type": "application/json", "X-Sync-Key": cfg["key"]},
    )
    try:
        with urllib.request.urlopen(req, timeout=120) as resp:
            return json.loads(resp.read().decode("utf-8"))
    except urllib.error.HTTPError as e:
        detail = ""
        try:
            detail = json.loads(e.read().decode("utf-8")).get("detail", "")
        except Exception:
            pass
        raise SyncError(f"الخادم رفض الطلب ({e.code}) {detail}")
    except urllib.error.URLError as e:
        raise SyncError(f"تعذّر الاتصال بالخادم: {e.reason}")


def send(cfg, rows, returns, dry_run, force=False):
    return call(cfg, "POST", "/api/v1/debt-sync/apply", {"rows": rows, "ignored_returns": returns, "dry_run": dry_run, "force": force})


def print_report(res):
    c, t = res["counts"], res["totals"]
    print(f"  قوائم صالحة بالملف: {c['file_rows_valid']}   مرتجعات متجاهلة: {c['ignored_returns']}")
    print(f"  ستُضاف: {c['to_insert']}   ستتحدّث: {c['to_update']}   ستُقفل: {c['to_close']}   بدون تغيير: {c['unchanged']}   مرفوضة: {c['rejected']}")
    print(f"  المتبقي قبل: {t['remaining_before']:,.2f}   بالملف: {t['remaining_in_file']:,.2f}   بعد التطبيق: {t['remaining_after']:,.2f}")
    if c.get("kept_for_pending_collections"):
        print(f"  تحصيلات المندوبين غير المستلمة (مخصومة من متبقي فينكس): {c['kept_for_pending_collections']} قائمة، {t['pending_collections_total']:,.2f}")
    if res.get("blocked_reason"):
        print("  ⚠ محجوز:", res["blocked_reason"])


def save_report(res, path):
    """Excel report: what changes and what was refused (customers to add / names to fix)."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "مرفوض"
    ws.append(["الصف", "الزبون", "رقم القائمة", "المبلغ", "الباقي", "التاريخ", "السبب", "اسم مشابه بالتطبيق"])
    for r in res["rejected"]:
        ws.append([r["row"], r["customer_name"], r["invoice_number"], r["amount"], r["remaining"], r["invoice_date"], r["reason"], r["similar"]])
    for title, key, cols in (
        ("تحصيلات غير مستلمة", "pending_collections", ["invoice_number", "customer_name", "program_remaining", "collected_in_app", "remaining"]),
        ("ستُضاف", "inserted", ["invoice_number", "customer_name", "amount", "remaining_amount"]),
        ("ستتحدّث", "updated", ["invoice_number", "customer_name", "remaining_from", "remaining_to", "amount_from", "amount_to"]),
        ("ستُقفل", "closed", ["invoice_number", "customer_name", "remaining_was"]),
    ):
        sh = wb.create_sheet(title)
        sh.append(cols)
        for r in res[key]:
            sh.append([r[c] for c in cols])
    wb.save(path)


def file_hash(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


# ---------------------------------------------------------------- commands
def cmd_check(path):
    rows, returns, info = read_export(path)
    print("الملف سليم:", json.dumps(info, ensure_ascii=False))


def cmd_send(path, dry_run, force=False):
    cfg = load_config()
    rows, returns, info = read_export(path)
    print("الملف:", json.dumps(info, ensure_ascii=False))
    res = send(cfg, rows, returns, dry_run, force)
    print_report(res)
    reports = os.path.join(HERE, "reports")  # never next to the export: the watcher reads the inbox
    os.makedirs(reports, exist_ok=True)
    stamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    out = os.path.join(reports, f"{stamp}_{'dry' if dry_run else 'apply'}.xlsx")
    save_report(res, out)
    print("  التقرير:", out)
    if not dry_run:
        print("  تم التطبيق، رقم المزامنة:", res["run_id"]) if res["applied"] else print("  لم يُطبَّق شيء")
    return res


def cmd_watch():
    sys.stdout.reconfigure(line_buffering=True)  # the log is readable while it runs
    cfg = load_config()
    inbox = cfg["inbox"]
    for sub in ("", "done", "failed", "held"):
        os.makedirs(os.path.join(inbox, sub), exist_ok=True)
    state_path = os.path.join(inbox, ".last_hash")
    last = open(state_path).read().strip() if os.path.exists(state_path) else ""
    print("أراقب:", inbox)
    while True:
        try:
            for name in sorted(os.listdir(inbox)):
                path = os.path.join(inbox, name)
                if not (os.path.isfile(path) and name.lower().endswith(".xlsx") and not name.startswith("~$")):
                    continue
                # wait until the export has finished being written
                size = (os.path.getsize(path), os.path.getmtime(path))
                time.sleep(5)
                if size != (os.path.getsize(path), os.path.getmtime(path)):
                    continue
                stamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
                h = file_hash(path)
                if h == last:
                    shutil.move(path, os.path.join(inbox, "done", f"{stamp}_duplicate_{name}"))
                    continue
                try:
                    rows, returns, info = read_export(path)
                    res = send(cfg, rows, returns, dry_run=False)
                    save_report(res, os.path.join(inbox, "done" if res["applied"] else "held", f"{stamp}_report.xlsx"))
                    dest = "done" if res["applied"] else "held"
                    print(stamp, "طُبّق" if res["applied"] else "محجوز:", res.get("blocked_reason") or "", res["counts"])
                    if res["applied"]:
                        last = h
                        open(state_path, "w").write(h)
                except FileBusy:
                    continue  # still open somewhere: look again on the next round
                except SyncError as e:
                    dest = "failed"
                    print(stamp, "فشل:", e)
                    with open(os.path.join(inbox, "failed", f"{stamp}_error.txt"), "w", encoding="utf-8") as f:
                        f.write(str(e))
                shutil.move(path, os.path.join(inbox, dest, f"{stamp}_{name}"))
        except Exception as e:  # keep watching whatever happens
            print("خطأ:", e)
        time.sleep(cfg["interval"])


def main(argv):
    if len(argv) < 2:
        print(__doc__)
        return 2
    cmd = argv[1]
    try:
        if cmd == "check" and len(argv) >= 3:
            cmd_check(argv[2])
        elif cmd == "dry" and len(argv) >= 3:
            cmd_send(argv[2], True)
        elif cmd == "apply" and len(argv) >= 3:
            cmd_send(argv[2], False, "--force" in argv)
        elif cmd == "undo" and len(argv) >= 3:
            cfg = load_config()
            print(call(cfg, "POST", f"/api/v1/debt-sync/undo/{int(argv[2])}"))
            # the same export must be applicable again after an undo
            last = os.path.join(cfg["inbox"], ".last_hash")
            if os.path.exists(last):
                os.remove(last)
        elif cmd == "runs":
            for r in call(load_config(), "GET", "/api/v1/debt-sync/runs"):
                print(r["id"], r["status"], r["created_at"], r.get("summary", {}).get("counts") if r.get("summary") else "")
        elif cmd == "watch":
            cmd_watch()
        else:
            print(__doc__)
            return 2
    except SyncError as e:
        print("خطأ:", e)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
